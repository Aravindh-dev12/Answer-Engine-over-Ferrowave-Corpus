import json
import csv
from pathlib import Path
from typing import Dict, Any, Optional
from bs4 import BeautifulSoup
import pypdf
import docx

class DocumentParser:
    @staticmethod
    def parse_file(file_path: Path) -> str:
        ext = file_path.suffix.lower()
        if ext in [".md", ".txt"]:
            return DocumentParser._parse_text(file_path)
        elif ext == ".html":
            return DocumentParser._parse_html(file_path)
        elif ext == ".json":
            return DocumentParser._parse_json(file_path)
        elif ext == ".csv":
            return DocumentParser._parse_csv(file_path)
        elif ext == ".pdf":
            return DocumentParser._parse_pdf(file_path)
        elif ext == ".docx":
            return DocumentParser._parse_docx(file_path)
        else:
            try:
                return file_path.read_text(encoding="utf-8", errors="replace")
            except Exception:
                return ""

    @staticmethod
    def _parse_text(path: Path) -> str:
        try:
            return path.read_text(encoding="utf-8", errors="replace")
        except Exception:
            return ""

    @staticmethod
    def _parse_html(path: Path) -> str:
        try:
            content = path.read_text(encoding="utf-8", errors="replace")
            soup = BeautifulSoup(content, "html.parser")
            # Preserve headings, paragraphs, table rows
            for tag in soup(["script", "style"]):
                tag.decompose()
            text = soup.get_text(separator="\n", strip=True)
            return text
        except Exception as e:
            return f"Error parsing HTML {path.name}: {e}"

    @staticmethod
    def _parse_json(path: Path) -> str:
        try:
            content = path.read_text(encoding="utf-8", errors="replace")
            data = json.loads(content)
            # Structured text representation
            lines = []
            if isinstance(data, dict):
                for k, v in data.items():
                    if isinstance(v, (dict, list)):
                        lines.append(f"## {k}")
                        lines.append(json.dumps(v, indent=2))
                    else:
                        lines.append(f"{k}: {v}")
            elif isinstance(data, list):
                for item in data:
                    lines.append(json.dumps(item, indent=2))
            return "\n\n".join(lines)
        except Exception as e:
            return f"Error parsing JSON {path.name}: {e}"

    @staticmethod
    def _parse_csv(path: Path) -> str:
        try:
            lines = []
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                headers = next(reader, None)
                if headers:
                    lines.append("Columns: " + " | ".join(headers))
                    for row in reader:
                        if not any(row):
                            continue
                        row_desc = [f"{h}: {v.strip()}" for h, v in zip(headers, row) if v.strip()]
                        lines.append(" - " + ", ".join(row_desc))
            return "\n".join(lines)
        except Exception as e:
            return f"Error parsing CSV {path.name}: {e}"

    @staticmethod
    def _parse_pdf(path: Path) -> str:
        try:
            reader = pypdf.PdfReader(str(path))
            pages = []
            for i, page in enumerate(reader.pages):
                txt = page.extract_text()
                if txt:
                    pages.append(f"--- Page {i+1} ---\n{txt}")
            return "\n\n".join(pages)
        except Exception as e:
            return f"Error parsing PDF {path.name}: {e}"

    @staticmethod
    def _parse_docx(path: Path) -> str:
        try:
            doc = docx.Document(str(path))
            lines = []
            for p in doc.paragraphs:
                txt = p.text.strip()
                if txt:
                    lines.append(txt)
            for table in doc.tables:
                for row in table.rows:
                    row_txt = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_txt:
                        lines.append(row_txt)
            return "\n\n".join(lines)
        except Exception as e:
            return f"Error parsing DOCX {path.name}: {e}"
